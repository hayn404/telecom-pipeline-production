import re
import io
from collections import defaultdict

import pandas as pd
import streamlit as st
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

PFX_RE    = re.compile(r"PFX=([^,]+)")
RSNAME_RE = re.compile(r'RSNAME="([^"]+)"')


def parse_cnacld_file(content: str) -> list[dict]:
    entries = []
    lines = content.splitlines()
    inside_block = False

    for line in lines:
        stripped = line.strip()

        # Detect CNACLD block start
        if re.search(r'\*\s*MOC\s*=\s*CNACLD\s*\*', stripped):
            inside_block = True
            continue

        # Detect next block — stop
        if inside_block and re.search(r'\*\s*MOC\s*=', stripped):
            break

        # Parse ADD CNACLD lines (also works if file is already pre-extracted)
        if stripped.startswith("ADD CNACLD"):
            inside_block = True  # treat pre-extracted files as already inside block
            pfx_m  = PFX_RE.search(stripped)
            rsn_m  = RSNAME_RE.search(stripped)
            if pfx_m and rsn_m:
                entries.append({
                    "PFX":    pfx_m.group(1).strip(),
                    "RSNAME": rsn_m.group(1).strip(),
                })

    return entries


def build_excel(df: pd.DataFrame) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CNACLD Entries"

    HEADER_FILL = PatternFill("solid", fgColor="2E4057")
    HEADER_FONT = Font(bold=True, color="FFFFFF")
    ALT_FILL    = PatternFill("solid", fgColor="F2F2F2")

    headers = list(df.columns)
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        ws.append(list(row))
        if row_idx % 2 == 0:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = ALT_FILL

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 35
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_cnacld_tab():
    st.header("📡 CNACLD Extractor")
    st.markdown("Extract all **CNACLD** entries (PFX → RSNAME) from an MSC MML file and export to Excel.")

    # --- Instructions box ---
    with st.expander("ℹ️ How to prepare your file before uploading", expanded=False):
        st.markdown("""
**The full MSC file can be 200MB+ which is too large to upload directly.**
First run this command on the company computer to extract only the CNACLD block:

**Windows PowerShell:**
```powershell
$src    = "C:\\path\\to\\ALLME_file.txt"
$output = "C:\\Users\\YourName\\Desktop\\MSC_cnacld.txt"

$capture = $false
$writer  = [System.IO.StreamWriter]::new($output)
foreach ($line in [System.IO.File]::ReadLines($src)) {
    if ($line -match '\\*\\s*MOC\\s*=\\s*CNACLD\\s*\\*') { $capture = $true; continue }
    if ($capture -and $line -match '\\*\\s*MOC\\s*=') { break }
    if ($capture -and $line.Trim() -ne '') { $writer.WriteLine($line) }
}
$writer.Close()
```
Then upload the small output file here.
        """)

    st.divider()

    uploaded_file = st.file_uploader(
        "Upload extracted CNACLD file (.txt or .sh)",
        type=["txt", "sh"],
        help="Upload the small file produced by the PowerShell extraction script"
    )

    if uploaded_file is None:
        st.info("Upload a file above to get started.")
        return

    try:
        content = uploaded_file.read().decode("utf-8", errors="ignore")
    except Exception as e:
        st.error(f"Could not read file: {e}")
        return

    entries = parse_cnacld_file(content)

    if not entries:
        st.warning("No CNACLD entries found in this file. Make sure the file contains ADD CNACLD lines.")
        return

    df = pd.DataFrame(entries)

    st.success(f"Found **{len(df)}** CNACLD entries from `{uploaded_file.name}`")

    # --- Stats row ---
    col1, col2 = st.columns(2)
    with col1:
        st.metric("Total Entries", len(df))
    with col2:
        st.metric("Unique RSNAMEs", df["RSNAME"].nunique())

    st.divider()

    # --- Filter ---
    search = st.text_input("🔍 Filter by PFX or RSNAME", placeholder="Type to search...")
    if search:
        mask = (
            df["PFX"].str.contains(search, case=False, na=False) |
            df["RSNAME"].str.contains(search, case=False, na=False)
        )
        df_display = df[mask]
        st.caption(f"Showing {len(df_display)} of {len(df)} entries")
    else:
        df_display = df

    st.dataframe(df_display, use_container_width=True, height=450)

    # --- Download ---
    st.divider()
    file_stem = uploaded_file.name.rsplit(".", 1)[0]
    excel_bytes = build_excel(df)

    st.download_button(
        label="📥 Download Excel",
        data=excel_bytes,
        file_name=f"{file_stem}_CNACLD.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary"
    )
