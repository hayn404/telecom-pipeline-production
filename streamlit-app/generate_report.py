#!/usr/bin/env python3
"""
Automated daily Excel report generator for e& CS Core Operations.

Called by daily_pipeline.sh after the ClickHouse sync step:
    docker exec telecom-prod-streamlit-app python /app/generate_report.py

Output: /app/reports/eand_core_stats_YYYY-MM-DD.xlsx  (one file per day)
        Host path: /mnt/raid5/production-deployment/reports/

Behaviour:
- Generates today's file on every run (data may have been refreshed).
- Skips files for past dates that already exist (past data is immutable).
- Auto-deletes files older than KEEP_DAYS (default 15).
"""

import io
import os
import sys
from datetime import date, timedelta
from pathlib import Path

import clickhouse_connect
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
REPORTS_DIR = Path(os.getenv("REPORTS_DIR", "/app/reports"))
KEEP_DAYS   = int(os.getenv("REPORT_KEEP_DAYS", "15"))
CH_HOST     = os.getenv("CLICKHOUSE_HOST", "telecom_prod_clickhouse")
CH_PORT     = int(os.getenv("CLICKHOUSE_PORT", "8123"))

EPS_5G_LIST = [
    "5115", "503", "554", "524", "585", "536", "530", "535",
    "511", "533", "537", "531", "540", "578", "579", "580",
    "592", "593", "594", "595", "557", "5111", "5113",
]

# ---------------------------------------------------------------------------
# ClickHouse helper
# ---------------------------------------------------------------------------
def get_client():
    return clickhouse_connect.get_client(
        host=CH_HOST, port=CH_PORT,
        username="default", password="",
        compress=True, connect_timeout=10, send_receive_timeout=60,
    )


def q(client, query):
    """Run a query and return result rows as a list of tuples."""
    return list(client.query(query).result_rows)


# ---------------------------------------------------------------------------
# Excel styling helper (mirrors _write_styled_sheet in app.py)
# ---------------------------------------------------------------------------
def _write_styled_sheet(ws, df, title=None):
    """Write a DataFrame to an openpyxl worksheet with e& brand styling."""
    HEADER_COLOR = "A40000"
    ALT_COLOR    = "FFF5F5"

    header_fill = PatternFill(start_color=HEADER_COLOR, end_color=HEADER_COLOR, fill_type="solid")
    alt_fill    = PatternFill(start_color=ALT_COLOR,    end_color=ALT_COLOR,    fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    data_font   = Font(size=10)
    thin_side   = Side(style="thin", color="DDDDDD")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    center      = Alignment(horizontal="center", vertical="center")

    if title:
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=13, color=HEADER_COLOR)
        ws.append([])

    header_row = ws.max_row + 1
    for ci, col_name in enumerate(df.columns, 1):
        c = ws.cell(row=header_row, column=ci, value=str(col_name))
        c.fill   = header_fill
        c.font   = header_font
        c.alignment = center
        c.border = cell_border

    for ri, row_vals in enumerate(df.itertuples(index=False), header_row + 1):
        fill = alt_fill if ri % 2 == 0 else PatternFill()
        for ci, val in enumerate(row_vals, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.fill   = fill
            c.font   = data_font
            c.alignment = center
            c.border = cell_border

    for ci, col_name in enumerate(df.columns, 1):
        col_vals = df.iloc[:, ci - 1].astype(str)
        max_len  = max(len(str(col_name)), col_vals.str.len().max() if len(df) > 0 else 0)
        ws.column_dimensions[get_column_letter(ci)].width = min(int(max_len) + 4, 45)


# ---------------------------------------------------------------------------
# Per-date Excel generation
# ---------------------------------------------------------------------------
def create_excel_for_date(client, date_str):
    """Query ClickHouse for date_str and return a formatted Excel file as bytes."""
    eps_5g_str = "', '".join(EPS_5G_LIST)
    safe_pct = lambda num, den: f"{num / den * 100:.1f}%" if den > 0 else "N/A"

    # --- Voice metrics ---
    rows = q(client, f"""
        SELECT
            uniqExact(MSISDN)                                                                          AS total_subs,
            uniqExactIf(MSISDN, TICK = '215')                                                         AS volte_subs,
            uniqExactIf(MSISDN, EpsAccessRestriction = '0')                                           AS vowifi_subs,
            uniqExactIf(MSISDN, length(EpsProfileId) >= 3
                AND (EpsProfileId LIKE '3%' OR EpsProfileId LIKE '5%')
                AND EpsProfileId = PDPCP)                                                              AS fiveg_subs,
            uniqExactIf(MSISDN, VLRADD != '' AND VLRADD IS NOT NULL
                AND NOT startsWith(VLRADD, '1920117900'))                                              AS roaming_subs,
            uniqExactIf(MSISDN, VLRADD != '' AND VLRADD IS NOT NULL
                AND startsWith(VLRADD, '1920117900'))                                                  AS local_subs,
            uniqExactIf(MSISDN, (VLRADD IS NULL OR VLRADD = '')
                AND (TICK IS NULL OR TICK = '') AND (PDPCP IS NULL OR PDPCP = '')
                AND length(MSISDN) >= 9 AND length(MSISDN) <= 11 AND MSISDN LIKE '%611%')             AS fvno_subs
        FROM default.dump
        WHERE CDRtime = '{date_str}'
    """)
    total_subs, volte_subs, vowifi_subs, fiveg_subs, roaming_subs, local_subs, fvno_subs = rows[0]

    # --- CS service distribution ---
    cs_rows = q(client, f"""
        SELECT
            CASE
                WHEN TICK = '201' THEN 'CS RPT'
                WHEN TICK = '205' THEN 'CS Call Screening'
                WHEN TICK = '203' THEN 'CS RPT + Call Screening'
                WHEN TICK = '190' AND (EpsProfileId = '' OR EpsProfileId IS NULL) THEN 'Normal CS Subscriber'
                WHEN (PDPCP = '' OR PDPCP IS NULL) AND TICK = '242' THEN 'PreActive Dial'
                ELSE NULL
            END AS service_type,
            uniqExact(MSISDN) AS subscribers
        FROM default.dump
        WHERE CDRtime = '{date_str}'
          AND TICK IS NOT NULL AND TICK != ''
          AND TICK != '215'
        GROUP BY service_type
        HAVING service_type IS NOT NULL
        ORDER BY subscribers DESC
    """)
    df_cs = pd.DataFrame(cs_rows, columns=["Service Type", "Subscribers"])
    total_cs = df_cs["Subscribers"].sum()
    df_cs["Percentage"] = df_cs["Subscribers"].apply(lambda x: safe_pct(x, total_cs))

    # --- Network distribution ---
    net_rows = q(client, f"""
        SELECT
            CASE
                WHEN length(EpsProfileId) >= 3
                    AND (EpsProfileId LIKE '3%' OR EpsProfileId LIKE '5%')
                    AND EpsProfileId = PDPCP THEN '5G'
                WHEN TICK = '215' THEN '4G VoLTE'
                WHEN TICK IN ('190','201','203','205')
                    AND (EpsProfileId = '' OR EpsProfileId IS NULL) THEN '2G/3G'
                WHEN TICK IN ('190','201','203','205')
                    AND (EpsProfileId != '' AND EpsProfileId IS NOT NULL) THEN '4G'
                ELSE NULL
            END AS network_type,
            uniqExact(MSISDN) AS subscribers
        FROM default.dump
        WHERE CDRtime = '{date_str}'
        GROUP BY network_type
        HAVING network_type IS NOT NULL
        ORDER BY subscribers DESC
    """)
    df_net = pd.DataFrame(net_rows, columns=["Network Type", "Subscribers"])
    total_net = df_net["Subscribers"].sum()
    df_net["% of Total"] = df_net["Subscribers"].apply(lambda x: safe_pct(x, total_net))

    # --- VoWiFi status ---
    wifi_rows = q(client, f"""
        SELECT
            CASE
                WHEN EpsAccessRestriction = '0' AND PDPCP != ''
                    AND PDPCP IN ('401','467','412','405','419','407','426','420') THEN '4G + WiFi'
                WHEN EpsAccessRestriction = '0' AND PDPCP != ''
                    AND PDPCP IN ('501','567','512','505','519','507','526','520') THEN '5G + WiFi'
                WHEN EpsAccessRestriction = '32' THEN 'Data Only (No VoWiFi)'
                WHEN EpsAccessRestriction = '63' THEN 'All Data Barred'
                ELSE NULL
            END AS status,
            uniqExact(MSISDN) AS subscribers
        FROM default.dump
        WHERE CDRtime = '{date_str}'
          AND EpsAccessRestriction IS NOT NULL AND EpsAccessRestriction != ''
        GROUP BY status
        HAVING status IS NOT NULL
        ORDER BY subscribers DESC
    """)
    df_wifi = pd.DataFrame(wifi_rows, columns=["VoWiFi Status", "Subscribers"])
    total_wifi = df_wifi["Subscribers"].sum()
    df_wifi["% of Total"] = df_wifi["Subscribers"].apply(lambda x: safe_pct(x, total_wifi))

    # --- Call features ---
    feat_rows = q(client, f"""
        SELECT
            uniqExactIf(MSISDN, CFU  = '1') AS cfu,
            uniqExactIf(MSISDN, CFB  = '1') AS cfb,
            uniqExactIf(MSISDN, HOLD = '1') AS hold,
            uniqExactIf(MSISDN, CAW  = '1') AS caw,
            uniqExactIf(MSISDN, DCF  = '1') AS mcn,
            uniqExactIf(MSISDN, CLIR = '1') AS clir,
            uniqExactIf(MSISDN, COLP = '1') AS colp,
            uniqExact(MSISDN)               AS total
        FROM default.dump
        WHERE CDRtime = '{date_str}'
    """)
    cfu, cfb, hold, caw, mcn, clir, colp, feat_total = feat_rows[0]
    df_feat = pd.DataFrame({
        "Feature":     ["CFU", "CFB", "HOLD", "CAW", "MCN", "CLIR", "COLP"],
        "Subscribers": [cfu, cfb, hold, caw, mcn, clir, colp],
        "% Adoption":  [safe_pct(v, feat_total) for v in [cfu, cfb, hold, caw, mcn, clir, colp]],
    })

    # --- Roaming by country ---
    country_rows = q(client, f"""
        SELECT
            CASE
                WHEN startsWith(VLRADD, '19966') THEN 'Saudi Arabia'
                WHEN startsWith(VLRADD, '19971') THEN 'UAE'
                WHEN startsWith(VLRADD, '19218') THEN 'Libya'
                WHEN startsWith(VLRADD, '19249') THEN 'Sudan'
                WHEN startsWith(VLRADD, '19965') THEN 'Kuwait'
                WHEN startsWith(VLRADD, '19962') THEN 'Jordan'
                WHEN startsWith(VLRADD, '19963') THEN 'Syria'
                WHEN startsWith(VLRADD, '19964') THEN 'Iraq'
                WHEN startsWith(VLRADD, '19968') THEN 'Oman'
                WHEN startsWith(VLRADD, '19974') THEN 'Qatar'
                WHEN startsWith(VLRADD, '19973') THEN 'Bahrain'
                WHEN startsWith(VLRADD, '19967') THEN 'Yemen'
                WHEN startsWith(VLRADD, '19961') THEN 'Lebanon'
                WHEN startsWith(VLRADD, '19972') THEN 'Palestine'
                WHEN startsWith(VLRADD, '19212') THEN 'Morocco'
                WHEN startsWith(VLRADD, '19213') THEN 'Algeria'
                WHEN startsWith(VLRADD, '19216') THEN 'Tunisia'
                WHEN startsWith(VLRADD, '19251') THEN 'Ethiopia'
                WHEN startsWith(VLRADD, '1944')  THEN 'UK'
                WHEN startsWith(VLRADD, '1933')  THEN 'France'
                WHEN startsWith(VLRADD, '1949')  THEN 'Germany'
                WHEN startsWith(VLRADD, '1939')  THEN 'Italy'
                WHEN startsWith(VLRADD, '1934')  THEN 'Spain'
                WHEN startsWith(VLRADD, '1931')  THEN 'Netherlands'
                ELSE 'Other'
            END AS country,
            uniq(MSISDN) AS subscriber_count
        FROM default.dump
        WHERE CDRtime = '{date_str}'
          AND VLRADD != '' AND VLRADD IS NOT NULL
          AND NOT startsWith(VLRADD, '1920117900')
        GROUP BY country
        HAVING country != 'Other'
        ORDER BY subscriber_count DESC
        LIMIT 10
    """)
    df_country = pd.DataFrame(country_rows, columns=["Country", "Subscribers"])
    if not df_country.empty:
        total_roam = df_country["Subscribers"].sum()
        df_country["% of Roaming"] = df_country["Subscribers"].apply(lambda x: safe_pct(x, total_roam))

    # --- Roaming by operator ---
    operator_rows = q(client, f"""
        SELECT
            CASE
                WHEN startsWith(VLRADD, '1996650')  THEN 'Saudi Arabia - STC'
                WHEN startsWith(VLRADD, '1996656')  THEN 'Saudi Arabia - Mobily'
                WHEN startsWith(VLRADD, '1996659')  THEN 'Saudi Arabia - Zain'
                WHEN startsWith(VLRADD, '1997150')  THEN 'UAE - Etisalat'
                WHEN startsWith(VLRADD, '1997155')  THEN 'UAE - du'
                WHEN startsWith(VLRADD, '19965500') THEN 'Kuwait - STC'
                WHEN startsWith(VLRADD, '1996596')  THEN 'Kuwait - Zain'
                WHEN startsWith(VLRADD, '199656')   THEN 'Kuwait - Ooredoo'
                WHEN startsWith(VLRADD, '1997477')  THEN 'Qatar - Ooredoo'
                WHEN startsWith(VLRADD, '1997455')  THEN 'Qatar - Vodafone'
                WHEN startsWith(VLRADD, '1996279')  THEN 'Jordan - Orange'
                WHEN startsWith(VLRADD, '1996277')  THEN 'Jordan - Umniah'
                WHEN startsWith(VLRADD, '1996278')  THEN 'Jordan - Zain'
                WHEN startsWith(VLRADD, '1997333')  THEN 'Bahrain - STC'
                WHEN startsWith(VLRADD, '1997336')  THEN 'Bahrain - Zain'
                WHEN startsWith(VLRADD, '1997339')  THEN 'Bahrain - Batelco'
                WHEN startsWith(VLRADD, '199647701') THEN 'Iraq - Asiacell'
                WHEN startsWith(VLRADD, '199647802') THEN 'Iraq - Korek'
                WHEN startsWith(VLRADD, '1996475')  THEN 'Iraq - Zain'
                WHEN startsWith(VLRADD, '1996394')  THEN 'Syria - Syriatel'
                WHEN startsWith(VLRADD, '1996393')  THEN 'Syria - MTN'
                WHEN startsWith(VLRADD, '1996771')  THEN 'Yemen - Sabafon'
                WHEN startsWith(VLRADD, '1996770')  THEN 'Yemen - MTN'
                WHEN startsWith(VLRADD, '1996895')  THEN 'Oman - Ooredoo'
                WHEN startsWith(VLRADD, '1996892')  THEN 'Oman - Omantel'
                WHEN startsWith(VLRADD, '1921891')  THEN 'Libya - Almadar'
                WHEN startsWith(VLRADD, '1921892')  THEN 'Libya - Libyana'
                WHEN startsWith(VLRADD, '1996134')  THEN 'Lebanon - Alfa'
                WHEN startsWith(VLRADD, '1996139')  THEN 'Lebanon - Touch'
                WHEN startsWith(VLRADD, '1997259')  THEN 'Palestine - Jawwal'
                WHEN startsWith(VLRADD, '1997256')  THEN 'Palestine - Ooredoo'
                WHEN startsWith(VLRADD, '1924991')  THEN 'Sudan - Zain'
                WHEN startsWith(VLRADD, '1924912')  THEN 'Sudan - Sudani'
                WHEN startsWith(VLRADD, '1924995')  THEN 'Sudan - Vivacell'
                WHEN startsWith(VLRADD, '192126639') THEN 'Morocco - Orange'
                WHEN startsWith(VLRADD, '19212661') THEN 'Morocco - Maroc Telecom'
                WHEN startsWith(VLRADD, '19212640') THEN 'Morocco - inwi'
                WHEN startsWith(VLRADD, '19213661') THEN 'Algeria - Mobilis'
                WHEN startsWith(VLRADD, '19213770') THEN 'Algeria - Djezzy'
                WHEN startsWith(VLRADD, '1921350')  THEN 'Algeria - Ooredoo'
                WHEN startsWith(VLRADD, '192165')   THEN 'Tunisia - Tunisie Telecom'
                WHEN startsWith(VLRADD, '1921622')  THEN 'Tunisia - Ooredoo'
                WHEN startsWith(VLRADD, '1921698')  THEN 'Tunisia - Orange'
                WHEN startsWith(VLRADD, '1925191')  THEN 'Ethiopia - Ethio Telecom'
                WHEN startsWith(VLRADD, '1990505')  THEN 'Turkey - Turk Telekom'
                WHEN startsWith(VLRADD, '1990559')  THEN 'Turkey - Turkcell'
                WHEN startsWith(VLRADD, '1990532')  THEN 'Turkey - Vodafone'
                WHEN startsWith(VLRADD, '19447781') THEN 'UK - O2'
                WHEN startsWith(VLRADD, '19447782') THEN 'UK - EE'
                WHEN startsWith(VLRADD, '19447802') THEN 'UK - EE'
                WHEN startsWith(VLRADD, '1944973')  THEN 'UK - Vodafone'
                WHEN startsWith(VLRADD, '19447953') THEN 'UK - O2'
                WHEN startsWith(VLRADD, '1944385')  THEN 'UK - Vodafone'
                WHEN startsWith(VLRADD, '194478297') THEN 'UK - Vodafone'
                WHEN startsWith(VLRADD, '19447624') THEN 'UK - Vodafone'
                WHEN startsWith(VLRADD, '19447797') THEN 'UK - Sure'
                WHEN startsWith(VLRADD, '1933660')  THEN 'France - Bouygues'
                WHEN startsWith(VLRADD, '1933689')  THEN 'France - Orange'
                WHEN startsWith(VLRADD, '1933609')  THEN 'France - SFR'
                WHEN startsWith(VLRADD, '1933695')  THEN 'France - Free'
                WHEN startsWith(VLRADD, '1949176')  THEN 'Germany - O2'
                WHEN startsWith(VLRADD, '1949177')  THEN 'Germany - Telekom'
                WHEN startsWith(VLRADD, '1949171')  THEN 'Germany - Vodafone'
                WHEN startsWith(VLRADD, '1949172')  THEN 'Germany - O2'
                WHEN startsWith(VLRADD, '1939391')  THEN 'Italy - Wind Tre'
                WHEN startsWith(VLRADD, '1939339')  THEN 'Italy - Wind Tre'
                WHEN startsWith(VLRADD, '1939320')  THEN 'Italy - Wind Tre'
                WHEN startsWith(VLRADD, '1939335')  THEN 'Italy - TIM'
                WHEN startsWith(VLRADD, '1939349')  THEN 'Italy - TIM'
                WHEN startsWith(VLRADD, '19393519') THEN 'Italy - TIM'
                WHEN startsWith(VLRADD, '19346404') THEN 'Spain - Airtel'
                WHEN startsWith(VLRADD, '1934607')  THEN 'Spain - Movistar'
                WHEN startsWith(VLRADD, '1934609')  THEN 'Spain - Orange'
                WHEN startsWith(VLRADD, '1934622')  THEN 'Spain - Yoigo'
                WHEN startsWith(VLRADD, '1992345')  THEN 'Pakistan - Jazz'
                WHEN startsWith(VLRADD, '1992333')  THEN 'Pakistan - Jazz'
                WHEN startsWith(VLRADD, '1992321')  THEN 'Pakistan - Jazz'
                WHEN startsWith(VLRADD, '199231')   THEN 'Pakistan - Jazz'
                WHEN startsWith(VLRADD, '1992300')  THEN 'Pakistan - Zong'
                WHEN startsWith(VLRADD, '199197')   THEN 'India - Aircell'
                WHEN startsWith(VLRADD, '199198')   THEN 'India - Airtel'
                WHEN startsWith(VLRADD, '1991981')  THEN 'India - Airtel'
                WHEN startsWith(VLRADD, '199195')   THEN 'India - Tata'
                WHEN startsWith(VLRADD, '199196')   THEN 'India - Vodafone Idea'
                WHEN startsWith(VLRADD, '193165')   THEN 'Netherlands - KPN'
                WHEN startsWith(VLRADD, '193162')   THEN 'Netherlands - Vodafone'
                WHEN startsWith(VLRADD, '193154')   THEN 'Netherlands - T-Mobile'
                WHEN startsWith(VLRADD, '193163')   THEN 'Netherlands - T-Mobile'
                ELSE ''
            END AS operator_label,
            uniq(MSISDN) AS subscriber_count
        FROM default.dump
        WHERE CDRtime = '{date_str}'
          AND VLRADD != '' AND VLRADD IS NOT NULL
          AND NOT startsWith(VLRADD, '1920117900')
        GROUP BY operator_label
        HAVING operator_label != ''
        ORDER BY subscriber_count DESC
        LIMIT 10
    """)
    df_op = pd.DataFrame(operator_rows, columns=["Operator", "Subscribers"])
    if not df_op.empty:
        split = df_op["Operator"].str.split(" - ", n=1, expand=True)
        df_op.insert(0, "Country", split[0])
        df_op["Operator"] = split[1] if 1 in split.columns else split[0]
        total_op = df_op["Subscribers"].sum()
        df_op["% of Shown"] = df_op["Subscribers"].apply(lambda x: safe_pct(x, total_op))

    # --- Data SIMs metrics ---
    ds_rows = q(client, f"""
        SELECT
            uniqExactIf(MSISDN,
                (TICK IS NULL OR TICK = '')
                AND ((PDPCP IS NOT NULL AND PDPCP != '') OR (EpsIndDefContextId IS NOT NULL AND EpsIndDefContextId != ''))
            ) AS data_sims_total,
            uniqExactIf(MSISDN,
                (TICK IS NULL OR TICK = '')
                AND ((PDPCP IS NOT NULL AND PDPCP != '') OR (EpsIndDefContextId IS NOT NULL AND EpsIndDefContextId != ''))
                AND (EpsProfileId IS NULL OR EpsProfileId = '')
            ) AS data_2g3g,
            uniqExactIf(MSISDN,
                (TICK IS NULL OR TICK = '')
                AND ((PDPCP IS NOT NULL AND PDPCP != '') OR (EpsIndDefContextId IS NOT NULL AND EpsIndDefContextId != ''))
                AND EpsProfileId IN ('{eps_5g_str}')
            ) AS data_5g,
            uniqExactIf(MSISDN,
                EpsIndDefContextId IS NOT NULL AND EpsIndDefContextId != ''
                AND NOT endsWith(EpsIndDefContextId, '37') AND NOT endsWith(EpsIndDefContextId, '00')
            ) AS corporate_data
        FROM default.dump
        WHERE CDRtime = '{date_str}'
    """)
    data_sims_total, data_2g3g, data_5g, corporate_data = ds_rows[0]
    data_4g = max(0, data_sims_total - data_2g3g - data_5g)

    # -----------------------------------------------------------------------
    # Build workbook
    # -----------------------------------------------------------------------
    wb = Workbook()
    wb.remove(wb.active)

    # Sheet 1: KPI Summary
    ws = wb.create_sheet("KPI Summary")
    kpi_rows = [
        ("Voice",     "Total Subscribers",      total_subs,      "100.0%"),
        ("Voice",     "VoLTE Subscribers",       volte_subs,      safe_pct(volte_subs,      total_subs)),
        ("Voice",     "VoWiFi Subscribers",      vowifi_subs,     safe_pct(vowifi_subs,      total_subs)),
        ("Voice",     "5G Subscribers",          fiveg_subs,      safe_pct(fiveg_subs,       total_subs)),
        ("Roaming",   "Roaming Subscribers",     roaming_subs,    safe_pct(roaming_subs,     total_subs)),
        ("Roaming",   "Egypt/Local Subscribers", local_subs,      safe_pct(local_subs,       total_subs)),
        ("Roaming",   "Fixed FVNO Subscribers",  fvno_subs,       safe_pct(fvno_subs,        total_subs)),
        ("Data SIMs", "Total Data SIMs",         data_sims_total, safe_pct(data_sims_total,  total_subs)),
        ("Data SIMs", "2G/3G Data SIMs",         data_2g3g,       safe_pct(data_2g3g,        data_sims_total)),
        ("Data SIMs", "4G Data SIMs",            data_4g,         safe_pct(data_4g,          data_sims_total)),
        ("Data SIMs", "5G Data SIMs",            data_5g,         safe_pct(data_5g,          data_sims_total)),
        ("Data SIMs", "Corporate Data",          corporate_data,  safe_pct(corporate_data,   total_subs)),
    ]
    _write_styled_sheet(ws, pd.DataFrame(kpi_rows, columns=["Category", "Metric", "Subscribers", "% of Total"]),
                        f"e& CS Core Operations — KPI Summary ({date_str})")

    # Sheet 2: CS Service Distribution
    if not df_cs.empty:
        ws = wb.create_sheet("CS Service Distribution")
        _write_styled_sheet(ws, df_cs, "CS Service Type Distribution")

    # Sheet 3: Network Distribution
    if not df_net.empty:
        ws = wb.create_sheet("Network Distribution")
        _write_styled_sheet(ws, df_net, "Network Technology Distribution")

    # Sheet 4: VoWiFi Status
    if not df_wifi.empty:
        ws = wb.create_sheet("VoWiFi Status")
        _write_styled_sheet(ws, df_wifi, "VoWiFi Access Distribution")

    # Sheet 5: Call Features
    ws = wb.create_sheet("Call Features")
    _write_styled_sheet(ws, df_feat, "Call Features Adoption")

    # Sheet 6: Roaming by Country
    if not df_country.empty:
        ws = wb.create_sheet("Roaming by Country")
        _write_styled_sheet(ws, df_country, "Top Roaming Countries")

    # Sheet 7: Roaming by Operator
    if not df_op.empty:
        ws = wb.create_sheet("Roaming by Operator")
        _write_styled_sheet(ws, df_op, "Top Roaming Operators")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    today_str = date.today().isoformat()

    print(f"[INFO] Report generator started — output dir: {REPORTS_DIR}")

    try:
        client = get_client()
    except Exception as exc:
        print(f"[ERR]  Cannot connect to ClickHouse: {exc}", file=sys.stderr)
        sys.exit(1)

    # Fetch up to KEEP_DAYS available dates from ClickHouse
    try:
        rows = client.query(
            f"SELECT DISTINCT CDRtime FROM default.dump ORDER BY CDRtime DESC LIMIT {KEEP_DAYS}"
        ).result_rows
    except Exception as exc:
        print(f"[ERR]  Failed to fetch available dates: {exc}", file=sys.stderr)
        sys.exit(1)

    available_dates = []
    for row in rows:
        d = row[0]
        available_dates.append(d.isoformat() if hasattr(d, "isoformat") else str(d))

    if not available_dates:
        print("[WARN] No dates found in ClickHouse dump table — nothing to export.")
        sys.exit(0)

    print(f"[INFO] Found {len(available_dates)} date(s): {', '.join(available_dates)}")

    # Generate reports
    for date_str in available_dates:
        output_path = REPORTS_DIR / f"eand_core_stats_{date_str}.xlsx"

        # Skip past dates that already have a file (data is immutable once loaded)
        if output_path.exists() and date_str != today_str:
            print(f"[SKIP] {date_str} — file exists, skipping")
            continue

        print(f"[GEN]  {date_str} — querying ClickHouse...")
        try:
            excel_bytes = create_excel_for_date(client, date_str)
            output_path.write_bytes(excel_bytes)
            size_kb = output_path.stat().st_size // 1024
            print(f"[OK]   {date_str} — saved {size_kb} KB → {output_path.name}")
        except Exception as exc:
            print(f"[ERR]  {date_str} — {exc}", file=sys.stderr)

    # Prune files older than KEEP_DAYS
    cutoff = date.today() - timedelta(days=KEEP_DAYS)
    for f in sorted(REPORTS_DIR.glob("eand_core_stats_*.xlsx")):
        try:
            file_date = date.fromisoformat(f.stem.replace("eand_core_stats_", ""))
            if file_date < cutoff:
                f.unlink()
                print(f"[DEL]  {f.name} — older than {KEEP_DAYS} days, removed")
        except ValueError:
            pass  # Skip files that don't match the naming pattern

    print("[INFO] Report generation complete.")


if __name__ == "__main__":
    main()
